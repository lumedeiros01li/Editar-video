"""Gera dados.json a partir da planilha (aba 'Retornos anuais'), com as mesmas fórmulas da aba 'Base 100 - 2000'.
Confere o resultado contra os valores calculados salvos na própria planilha."""
import hashlib, json, sys
from datetime import date
import openpyxl

ARQ = 'fonte-dados/comparacao_ativos_25_anos.xlsx'
wf = openpyxl.load_workbook(ARQ, data_only=False)
wv = openpyxl.load_workbook(ARQ, data_only=True)
ret = wv['Retornos anuais']

# colunas da aba 'Retornos anuais'
COL = {'ano': 1, 'POU': 2, 'CDI': 3, 'IPCA': 4, 'USD': 5, 'IBOV': 6, 'SPX_USD': 7, 'OURO_USD': 8}
linhas = []
for r in range(2, ret.max_row + 1):
    ano = ret.cell(r, 1).value
    if ano is None or not str(ano)[:4].isdigit():
        continue
    linhas.append({k: ret.cell(r, c).value for k, c in COL.items()})

# base 100 em 31/12/2000: a linha de 2000 só define a base, o primeiro retorno usado é o de 2001
ids = ['POU', 'CDI', 'OURO', 'IBOV', 'SPX', 'IPCA', 'USD']
serie = {k: [100.0] for k in ids}
retornos = {k: [] for k in ids}
anos = [str(linhas[0]['ano'])]
for l in linhas[1:]:
    r = {
        'POU': l['POU'], 'CDI': l['CDI'], 'IPCA': l['IPCA'], 'USD': l['USD'], 'IBOV': l['IBOV'],
        'SPX': (1 + l['SPX_USD']) * (1 + l['USD']) - 1,
        'OURO': (1 + l['OURO_USD']) * (1 + l['USD']) - 1,
    }
    for k in ids:
        retornos[k].append(r[k])
        serie[k].append(serie[k][-1] * (1 + r[k]))
    anos.append(str(l['ano']))

# conferência contra os valores calculados da aba 'Base 100 - 2000' (colunas B..H)
base = wv['Base 100 - 2000']
mapa = {'POU': 2, 'CDI': 3, 'OURO': 4, 'IBOV': 5, 'SPX': 6, 'USD': 7, 'IPCA': 8}
pior = 0
for i in range(len(anos)):
    for k, c in mapa.items():
        v = base.cell(2 + i, c).value
        pior = max(pior, abs(v - serie[k][i]))
if pior > 1e-6:
    sys.exit(f'diverge da planilha: {pior}')

resumo = wv['Resumo']
anos_periodo = resumo['B5'].value
# posição de cada ponto no eixo do tempo (anos desde 31/12/2000); 2026* é parcial
pos = [float(i) for i in range(len(anos) - 1)] + [len(anos) - 2 + (anos_periodo - (len(anos) - 2))]

fontes = []
fn = wv['Fontes e notas']
for r in range(2, fn.max_row + 1):
    if fn.cell(r, 1).value:
        fontes.append({'serie': fn.cell(r, 1).value, 'nota': fn.cell(r, 2).value})

nomes = {
    'POU': ('Poupança', 'POUPANÇA'), 'CDI': ('CDI (Selic)', 'CDI'), 'OURO': ('Ouro (R$)', 'OURO'),
    'IBOV': ('Ibovespa', 'IBOVESPA*'), 'SPX': ('S&P 500 (R$)', 'S&P 500*'), 'IPCA': ('IPCA (inflação)', 'INFLAÇÃO'),
    'USD': ('Dólar (Ptax)', 'DÓLAR'),
}
dados = {
    'titulo': 'Quanto R$ 100 viraram de 31/12/2000 até 30/09/2026 (valores brutos, sem imposto)',
    'base': {'data': '31/12/2000', 'valor': 100, 'ano': int(anos[0])},
    'fim': {'data': '30/09/2026', 'rotulo': 'set/2026'},
    'anos_periodo': anos_periodo,
    'anos': [a.rstrip('*') for a in anos],
    'posicao': pos,
    'series': {k: {'nome': nomes[k][0], 'rotulo': nomes[k][1], 'valores': [round(v, 6) for v in serie[k]],
                   'retornos': [round(v, 6) for v in retornos[k]]} for k in ids},
    'fontes': fontes,
    'arquivo': {'nome': 'comparacao_ativos_25_anos.xlsx', 'sha256': hashlib.sha256(open(ARQ, 'rb').read()).hexdigest(),
                'recebido_em': str(date(2026, 10, 24))},
    'conferencia': f'séries recalculadas batem com a aba Base 100 - 2000 (diferença máxima {pior:.2e})',
}
json.dump(dados, open('dados.json', 'w'), ensure_ascii=False, indent=1)
for k in ids:
    print(f"{k:5} {serie[k][-1]:8.1f}")
print('posições', pos[-3:], 'conferência', pior)
